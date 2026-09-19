"""Build docs/Navaid_Architecture.docx — assignment scoring / tradeoffs / AI write-up."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Navaid_Architecture.docx"
NAVY = RGBColor(0x12, 0x1A, 0x24)
BRASS = RGBColor(0x6F, 0x55, 0x20)
INK = RGBColor(0x16, 0x13, 0x0F)
SOFT = RGBColor(0x4D, 0x47, 0x3E)

WEIGHTS = {
    "Demand pressure": 0.22,
    "Congestion": 0.18,
    "Landside saturation": 0.18,
    "Unmet demand": 0.15,
    "Yield mix": 0.12,
    "Growth outlook": 0.10,
    "Capital feasibility": 0.05,
}
CONSTRAINTS = {
    "Landside": 1.00,
    "Mixed": 0.75,
    "Airside": 0.40,
    "Demand-bound": 0.30,
}
WATERFALL = [
    ("Demand pressure", 9.5),
    ("Congestion", 10.4),
    ("Landside saturation", 11.6),
    ("Unmet demand", 14.4),
    ("Yield mix", 8.8),
    ("Growth outlook", 6.5),
]


def _set_run(run, *, size=11, bold=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def _heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY if level == 1 else BRASS
        run.font.name = "Calibri"
    return p


def _para(doc, text, *, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    _set_run(run, size=size, color=INK)
    run.italic = italic
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    _set_run(run, size=11)
    return p


def _table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        _set_run(run, size=10, bold=True, color=NAVY)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = ""
            run = table.rows[r].cells[c].paragraphs[0].add_run(str(value))
            _set_run(run, size=10)
    doc.add_paragraph()
    return table


def _chart_to_stream(fig) -> BytesIO:
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def _weights_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    labels = list(WEIGHTS)
    values = list(WEIGHTS.values())
    colors = ["#121a24", "#1c2836", "#5c4634", "#1f4d3a", "#6f5520", "#3d4554", "#7a3226"]
    ax.barh(labels[::-1], values[::-1], color=colors[::-1])
    ax.set_xlabel("Weight (sums to 1.00)")
    ax.set_title("TEOI feature weights")
    ax.set_xlim(0, 0.28)
    for y, v in enumerate(values[::-1]):
        ax.text(v + 0.004, y, f"{v:.2f}", va="center", fontsize=8)
    fig.tight_layout()
    return _chart_to_stream(fig)


def _constraint_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.4, 3.2))
    labels = list(CONSTRAINTS)
    values = list(CONSTRAINTS.values())
    colors = ["#1f4d3a", "#5c4634", "#7a3226", "#3d4554"]
    bars = ax.bar(labels, values, color=colors)
    ax.set_ylim(0, 1.15)
    ax.set_ylabel("Multiplier after the weighted sum")
    ax.set_title("Constraint multipliers — the investment thesis in arithmetic")
    for bar, v in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, v + 0.03, f"{v:.2f}", ha="center", fontsize=9)
    fig.tight_layout()
    return _chart_to_stream(fig)


def _waterfall_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.4, 3.4))
    labels = [n for n, _ in WATERFALL]
    values = [v for _, v in WATERFALL]
    ax.bar(labels, values, color="#121a24")
    ax.set_ylabel("Contribution to TEOI")
    ax.set_title("Illustrative BDL waterfall (capital_feasibility dropped)")
    ax.tick_params(axis="x", rotation=20)
    ax.axhline(61.2, color="#6f5520", linestyle="--", linewidth=1, label="Weighted sum 61.2 × landside 1.00")
    ax.legend(fontsize=8)
    fig.tight_layout()
    return _chart_to_stream(fig)


def _ai_split_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.4, 3.2))
    labels = [
        "Ingest / warehouse",
        "TEOI / unmet /\nlong-haul / congestion",
        "Reconstruct &\ndecompose",
        "Narrate locked\nprose",
        "Number lock",
    ]
    engines = [1, 1, 0.35, 0.15, 1]
    gemini = [0, 0, 0.65, 0.85, 0]
    x = range(len(labels))
    ax.bar(x, engines, color="#121a24", label="Deterministic Python")
    ax.bar(x, gemini, bottom=engines, color="#c4a35a", label="Gemini")
    ax.set_xticks(list(x), labels, fontsize=8)
    ax.set_ylim(0, 1.25)
    ax.set_ylabel("Who owns the step")
    ax.set_title("Where AI is used — Gemini never ranks")
    ax.legend(frameon=False)
    fig.tight_layout()
    return _chart_to_stream(fig)


def _pipeline_chart() -> BytesIO:
    fig, ax = plt.subplots(figsize=(7.2, 2.4))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2)
    ax.axis("off")
    boxes = [
        (0.2, "Public\nfiles"),
        (2.0, "DuckDB\nwarehouse"),
        (3.8, "Engines\n+ traces"),
        (5.6, "Gemini\nnarration"),
        (7.4, "Number\nlock"),
        (9.0, "Answer +\nenvelope"),
    ]
    for x, label in boxes:
        ax.add_patch(
            plt.Rectangle((x, 0.55), 1.4, 1.05, fill=True, facecolor="#f3eee3", edgecolor="#121a24", lw=1.2)
        )
        ax.text(x + 0.7, 1.08, label, ha="center", va="center", fontsize=8, color="#121a24")
    for x in (1.6, 3.4, 5.2, 7.0, 8.8):
        ax.annotate("", xy=(x + 0.2, 1.08), xytext=(x - 0.2, 1.08), arrowprops=dict(arrowstyle="->", color="#6f5520"))
    ax.set_title("Navaid pipeline", loc="left", fontsize=11, color="#121a24")
    fig.tight_layout()
    return _chart_to_stream(fig)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Navaid")
    _set_run(run, size=28, bold=True, color=NAVY, name="Calibri")
    kicker = doc.add_paragraph()
    run = kicker.add_run("Airport Investment Intelligence Agent  ·  Architecture")
    _set_run(run, size=12, color=BRASS)

    _para(
        doc,
        "Assignment deliverable: scoring methodology, key tradeoffs, and where Gemini is used "
        "versus deterministic engines. Companion to docs/ARCHITECTURE.md in the repository.",
        italic=True,
    )
    _para(
        doc,
        "Public aviation files are ingested into a local DuckDB warehouse. Deterministic engines "
        "compute TEOI (with a full ScoringTrace), unmet demand, long-haul share, and congestion "
        "compares. A Gemini agent reconstructs follow-ups, decomposes multi-part questions, then "
        "may only call those engines as tools. A number lock rejects any digit that was not in a "
        "tool payload. Gemini may explain a rank; it may never produce one.",
    )

    doc.add_picture(_pipeline_chart(), width=Inches(6.4))
    cap = doc.add_paragraph()
    run = cap.add_run("Figure 1. End-to-end pipeline. Live FAA status skips the warehouse and overlays operations delay only.")
    _set_run(run, size=9, color=SOFT)
    cap.runs[0].italic = True

    _heading(doc, "1. Scoring methodology", 1)
    _heading(doc, "Investment thesis", 2)
    _para(
        doc,
        "The firm invests in terminal modernization. Profit is a capacity-unlock proxy — not PFC, "
        "bond, or NPV. Busy is not investable: an airport can already be maxed on slots, curfew, "
        "or a weak catchment. TEOI (Terminal Expansion Opportunity Index) is a 0–100 peer-relative "
        "score inside the comparison set S that the analyst named.",
    )
    _bullet(doc, "Landside × 1.00 — gates, holdrooms, security, bag claim, curb. Terminal capex can raise throughput.")
    _bullet(doc, "Mixed × 0.75 — landside and airside both bind. A terminal project only partially unlocks capacity.")
    _bullet(doc, "Airside × 0.40 — runways, slots, weather, ATC, noise curfew. More terminal does not create slots.")
    _bullet(doc, "Demand-bound × 0.30 — weak catchment, low load factor, or leakage already served nearby.")

    _heading(doc, "Peer-relative min-max", 2)
    _para(
        doc,
        "Each feature is stretched 0–1 inside S. Highest in S = 1, lowest = 0. Ask New England and "
        "BOS is the high end. Ask only BOS and every feature becomes 0.5, so mixed × 0.75 can print "
        "~37. That is not a new formula — it is a one-airport peer set. Follow-ups such as "
        "“why does BOS have this score?” reuse the last set; they must not re-rank BOS alone.",
    )

    _heading(doc, "Formula", 2)
    _para(
        doc,
        "TEOI = 100 × (0.22·demand_pressure + 0.18·congestion + 0.18·landside_saturation + "
        "0.15·unmet_demand + 0.12·yield_mix + 0.10·growth_outlook + 0.05·capital_feasibility) "
        "× constraint_multiplier.",
        italic=True,
    )
    doc.add_picture(_weights_chart(), width=Inches(6.1))
    cap = doc.add_paragraph()
    run = cap.add_run("Figure 2. Default TEOI weights in navaid/scoring/weights.py. They sum to 1.00.")
    _set_run(run, size=9, color=SOFT)
    cap.runs[0].italic = True

    _table(
        doc,
        ["Feature", "Weight", "Role"],
        [
            ["Demand pressure", "0.22", "Enplanements / YoY vs peers in S"],
            ["Congestion", "0.18", "Delay / cancel / utilization vs peers"],
            ["Landside saturation", "0.18", "Passengers per gate vs peers"],
            ["Unmet demand", "0.15", "Load-factor gap + TAF gap + leakage"],
            ["Yield mix", "0.12", "Long-haul / international mix as a yield proxy"],
            ["Growth outlook", "0.10", "TAF / CAGR outlook vs peers"],
            ["Capital feasibility", "0.05", "NPIAS/AIP; often the dropped weight"],
        ],
    )

    doc.add_picture(_constraint_chart(), width=Inches(6.1))
    cap = doc.add_paragraph()
    run = cap.add_run("Figure 3. Constraint multipliers applied after the weighted sum.")
    _set_run(run, size=9, color=SOFT)
    cap.runs[0].italic = True

    _heading(doc, "Drop-and-renormalize", 2)
    _para(
        doc,
        "If a feature is missing (for example no NPIAS for capital feasibility), that weight is "
        "removed and the remainder is rescaled to 1.0. Navaid never invents gate counts or TAF. "
        "The drop is listed on the ScoringTrace (weights_dropped, weights_used) and in the envelope. "
        "Example: drop 0.05; remaining mass 0.95; demand pressure used weight becomes 0.22 / 0.95 ≈ 0.232. "
        "Contributions are 100 × weight_used × scaled_0_1.",
    )

    doc.add_picture(_waterfall_chart(), width=Inches(6.1))
    cap = doc.add_paragraph()
    run = cap.add_run(
        "Figure 4. Illustrative BDL waterfall after drop-and-renormalize. Live ranks come from the warehouse, not this picture."
    )
    _set_run(run, size=9, color=SOFT)
    cap.runs[0].italic = True

    _heading(doc, "Unmet demand", 2)
    _bullet(doc, "Load-factor gap: served × max(0, LF − 0.85) / 0.85.")
    _bullet(doc, "Forecast gap: max(0, TAF_10y − implied_current_capacity) if TAF exists.")
    _bullet(doc, "Leakage: same-CBSA airports growing faster, only while origin load factor is at or above 85%.")
    _para(
        doc,
        "If load factor is below the seat-pressure rule, leakage is qualitative (for SFO: OAK/SJC) "
        "and not in the number. A TAF 10-year gap can still produce unmet passengers. Slot/curfew "
        "remain the investment caveat: a forecast gap is not a missing concourse.",
    )

    _heading(doc, "Long-haul", 2)
    _para(
        doc,
        "Great-circle kilometres on BTS T-100 segments, default Eurocontrol threshold 4000 km. "
        "The engine reports flight-segment share and passenger share, plus international share and "
        "optional percent over 6 hours. Cargo is excluded unless asked. Pin ANC→JFK ≈ 5420 km / 3370 mi. "
        "Warehouse T-100 is a filtered commercial sample, not every cargo and connecting itinerary. "
        "ANC→SEA→JFK is two segments, not true origin-and-destination.",
    )

    _heading(doc, "Congestion and ranker", 2)
    _para(
        doc,
        "Congestion is compared axis by axis: delay share, average arrival delay, cancellations, "
        "operations per runway, live FAA status, constraint type, and curfew text. There is no single "
        "congestion score. Shared ranks are 3, 3, 5; leftover ties break by ICAO ascending. Deterministic.",
    )

    _heading(doc, "2. Key tradeoffs", 1)
    _table(
        doc,
        ["Choice", "What we shipped", "What we refused", "Why"],
        [
            [
                "KPI",
                "Peer-relative TEOI + traces",
                "LLM ranking; national grade",
                "The brief requires deterministic scoring and a visible working.",
            ],
            [
                "Warehouse",
                "One DuckDB file",
                "Postgres / Snowflake / Pandas-as-SoR",
                "T-100 needs SQL; the assignment is local-first.",
            ],
            [
                "RAG",
                "DuckDB FTS, airport-keyed notes",
                "Chroma, embeddings, RAG-as-ranker",
                "Tens of notes. T-100 wins facts.",
            ],
            [
                "Agent loop",
                "Owned orchestrator",
                "LangChain, CrewAI, ADK, AFC",
                "Reconstruct, traces, and number lock must be visible Python.",
            ],
            [
                "Model",
                "gemini-2.5-flash",
                "Pro as default; OpenAI",
                "Numbers already come from engines. Flash is enough to narrate.",
            ],
            [
                "Profit",
                "Capacity-unlock proxy",
                "PFC / NPV / airline equity",
                "Honest scope. AAL is refused.",
            ],
            [
                "Long-haul",
                "4000 km segments",
                "True O&D; mixing km and miles",
                "T-100 is legs. ANC→JFK pins the unit.",
            ],
            [
                "Congestion",
                "Axis-by-axis + curfew",
                "Fake composite score",
                "LAX delay volume vs SNA curfew are different facts.",
            ],
        ],
    )
    _para(
        doc,
        "Stale warehouse snapshot (>90 days) is always listed as an uncertainty and pulls confidence "
        "to low. Missing required inputs also force low confidence. That is an honesty tradeoff: the "
        "demo will not pretend 2024 boardings are live 2026 operations.",
    )

    _heading(doc, "3. Where and how AI is used", 1)
    doc.add_picture(_ai_split_chart(), width=Inches(6.1))
    cap = doc.add_paragraph()
    run = cap.add_run("Figure 5. Gemini may reconstruct, decompose, and narrate. Engines own every number.")
    _set_run(run, size=9, color=SOFT)
    cap.runs[0].italic = True

    _para(doc, "Gemini is allowed to:")
    _bullet(doc, "Rewrite follow-ups (“those two”, “why is #2 above #3?”, “add PWM”) into a standalone reconstructed_query.")
    _bullet(doc, "Help decompose a compound question into closed intents — rules still emit every subgoal.")
    _bullet(doc, "Call a closed tool list: resolve_airport, rank_expansion, compare_congestion, longhaul_share, unmet_demand, airport_metrics, explain_teoi, search_corpus, get_live_status.")
    _bullet(doc, "Narrate locked traces in Markdown. Automatic function calling is off so the number lock can intercept every digit.")
    _bullet(doc, "Optional TTS on the workbench. Same vendor as the chat model.")

    _para(doc, "Gemini is forbidden to:")
    _bullet(doc, "Produce a TEOI, rank, percent, or gate count that is not in tool JSON.")
    _bullet(doc, "Re-rank a singleton when session traces already include a peer set.")
    _bullet(doc, "Treat RAG notes as a ranker. If notes and T-100 disagree, T-100 wins.")
    _bullet(doc, "Advise on securities (buy AAL), restaurants, or non-airport questions — those parts are refused, the rest still answers.")

    _heading(doc, "Every /ask", 2)
    _para(
        doc,
        "User question → reconstruct → decompose → plan tools → run engines → number lock + TEOI "
        "traces → narrate one section per subgoal → store session. The Answer object always carries "
        "assumptions, uncertainties, out-of-scope, confidence, sources, and as_of. Gradio and the "
        "designed site are skins over the same FastAPI contract.",
    )

    _heading(doc, "Canonical questions", 2)
    _bullet(doc, "New England expansion — peer-relative TEOI; BOS is scale; regionals can win landside; PWM is Portland Maine, not PDX.")
    _bullet(doc, "LA vs Santa Ana congestion — LA → LAX not the metro; Santa Ana → SNA not SAT; no composite score.")
    _bullet(doc, "Anchorage long-haul — 4000 km segments; flight share and passenger share; cargo out unless asked.")
    _bullet(doc, "SFO unmet — clamped formula; leakage gated on 85% load factor; slot/curfew still bind.")

    foot = doc.add_paragraph()
    run = foot.add_run(
        "Related repository files: docs/ARCHITECTURE.md (canonical markdown), docs/GLOSSARY.md, "
        "docs/DECISION_LOG.md. Default model gemini-2.5-flash. RAG is DuckDB FTS — no embedding model, no Chroma."
    )
    _set_run(run, size=9, color=SOFT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(OUT.name)
